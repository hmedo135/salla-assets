import openpyxl, shutil
SRC='/mnt/user-data/uploads/Downloads/'
OUT='/mnt/user-data/outputs/'
MP='[marketplace_id=A17E79C6D8DWNP]'
GH='https://raw.githubusercontent.com/hmedo135/salla-assets/main/amazon-listing/'
S='https://cdn.salla.sa/bXpaY/'
images=[S+'ylkKaXYmOlAuYZ6XROK69MWKLTpvgtDOKRJYOr6k.jpg',   # main: white bg + cable
        GH+'cover-nologo-1600.jpg',                          # cover (logo removed)
        S+'jjg6TfHsxSjxa9KbVihD8gLD60XPCbA461W8hrVx.png',   # school scene, no logo
        GH+'how-it-works-nologo-1600.jpg',                   # how it works, no logo
        S+'onv676qSTkMJKrF8QDC5eblVVGM87jpHsjmZVT9c.jpg',   # white bg antenna
        S+'WGpIm38bEshU00ejjhBZfnMUVYqnXi9fiO3yO7Zh.jpg',   # hospital
        S+'6YjmPIDWVmqcDawSPO1VGnvvAnL0ibbwssnovjo8.jpg']   # nurse

en_desc=("Press the remote from where you are and the alert reaches the watch of the person you are calling instantly - no SIM, no internet, no subscriptions.<br><br>"
 "<b>Perfect for school pick-up</b><br>Parents don't need to shout or get out of the car. As soon as you arrive, press the remote and the alert goes off on the student's watch, so they head straight to the car. No crowding at the gate and no waiting.<br><br>"
 "<b>Also suitable for hospitals</b><br>The patient presses the remote to call the nurse or companion wearing the watch.<br><br>"
 "<b>Specifications</b><br>- Transmission range: up to 400 meters in open areas<br>- Battery: up to 1 month in economy mode, up to 1 week in continuous mode<br>- Charges from any USB port<br>- Works without SIM, internet or monthly subscriptions<br><br>"
 "<b>Package contents</b><br>- Remote<br>- Watch<br>- USB charging cable")
en_bullets=["INSTANT CALL: Press the remote button and the alert rings instantly on the watch of the person you are calling",
 "SCHOOL PICK-UP: Call your child from the car without shouting, getting out or waiting at the school gate",
 "NO SIM, NO INTERNET: Works without a SIM card, internet connection or monthly subscriptions",
 "400 M RANGE & LONG BATTERY: Up to 400 meters in open areas; battery lasts up to 1 month in economy mode or 1 week in continuous mode; charges from any USB port",
 "PACKAGE CONTENTS: 1 remote, 1 watch and 1 USB charging cable - also suitable for hospitals to call a nurse or companion"]
en_kw="pager watch school pickup caller wireless calling system nurse call caregiver pager remote alarm watch student caller"

ar_title="ساعة إنذار ونداء لا سلكية عن بعد دون الحاجة إلى شريحة أو إتصال بالإنترنت للمدارس والمستشفيات"
ar_desc=("اضغط الريموت وأنت في مكانك، وتوصل الإشارة لساعة الشخص اللي تناديه فورًا، بدون شريحة ولا إنترنت ولا اشتراكات.<br><br>"
 "<b>مثالية لانصراف المدارس</b><br>ولي الأمر ما يحتاج ينادي بصوته ولا ينزل من سيارته. أول ما يوصل، يضغط الريموت ويطلع الإنذار في ساعة الطالب، فيروح لسيارة أهله مباشرة. بدون زحام عند البوابة وبدون انتظار.<br><br>"
 "<b>مناسبة للمستشفيات</b><br>المريض يضغط الريموت فيصل التنبيه لساعة الممرض أو المرافق.<br><br>"
 "<b>المواصفات</b><br>- مدى الإرسال: 400 متر في المساحات المفتوحة<br>- البطارية: الوضع الاقتصادي حتى شهر، ووضع التشغيل الدائم حتى أسبوع<br>- تُشحن من أي منفذ USB<br>- تعمل بدون شريحة وبدون إنترنت وبدون اشتراكات شهرية<br><br>"
 "<b>محتويات العبوة</b><br>- ريموت<br>- ساعة<br>- سلك شحن USB")
ar_bullets=["نداء فوري: اضغط زر الريموت وأنت في مكانك، ويرن الإنذار فورًا في ساعة الشخص اللي تناديه",
 "مثالية لانصراف المدارس: ولي الأمر ينادي الطالب من سيارته بدون ما ينزل أو يرفع صوته، وبدون زحام عند البوابة",
 "بدون شريحة وبدون إنترنت: تعمل بدون شريحة ولا اتصال إنترنت ولا اشتراكات شهرية",
 "مدى 400 متر وبطارية تدوم: مدى الإرسال حتى 400 متر في المساحات المفتوحة، والبطارية تدوم حتى شهر في الوضع الاقتصادي وحتى أسبوع في التشغيل الدائم، وتُشحن من أي منفذ USB",
 "محتويات العبوة: ريموت + ساعة + سلك شحن USB، ومناسبة أيضًا للمستشفيات لنداء الممرض أو المرافق"]
ar_kw="ساعة نداء ساعة مناداة الطلاب ساعة المدرسة جهاز نداء لاسلكي ريموت نداء نداء المريض جرس لاسلكي ساعة ناديني"

def build(src,dst,sheet,lang,vals):
    shutil.copy(src,dst)
    wb=openpyxl.load_workbook(dst,keep_vba=True)
    ws=wb[sheet]
    attrs={ws.cell(5,c).value:c for c in range(1,ws.max_column+1) if ws.cell(5,c).value}
    for k,v in vals.items():
        if k not in attrs: raise SystemExit('missing column '+k)
        ws.cell(7,attrs[k]).value=v
    wb.save(dst)
    print('saved',dst,len(vals),'fields')

L=lambda a,lang,n=1:f'{a}{MP}[language_tag={lang}]#{n}.value'
en={'contribution_sku#1.value':'NASCART-PAGER-01','product_type#1.value':'CONSUMER_ELECTRONICS','::record_action':'Edit (Partial Update)',
    f'model_number{MP}#1.value':'NASCART001', f'part_number{MP}#1.value':'NASCART001',
    f'main_product_image_locator{MP}#1.media_location':images[0],
    L('product_description','en_AE'):en_desc, L('generic_keyword','en_AE'):en_kw}
for i,u in enumerate(images[1:],1): en[f'other_product_image_locator_{i}{MP}#1.media_location']=u
for i,b in enumerate(en_bullets,1): en[L('bullet_point','en_AE',i)]=b
ar={'contribution_sku#1.value':'NASCART-PAGER-01','product_type#1.value':'CONSUMER_ELECTRONICS','::record_action':'تحرير (تحديث جزئي)',
    L('item_name','ar_AE'):ar_title, L('product_description','ar_AE'):ar_desc, L('generic_keyword','ar_AE'):ar_kw,
    L('color','ar_AE'):'أسود', L('included_components','ar_AE',1):'ريموت', L('included_components','ar_AE',2):'ساعة',
    L('included_components','ar_AE',3):'سلك شحن USB', L('warranty_description','ar_AE'):'3 أشهر'}
for i,b in enumerate(ar_bullets,1): ar[L('bullet_point','ar_AE',i)]=b
for b in en_bullets+ar_bullets: assert len(b)<=255, (len(b),b)
print('kw bytes', len(en_kw.encode()), len(ar_kw.encode()))
build(SRC+'CONSUMER_ELECTRONICS.xlsm',OUT+'NasCart-Amazon-EN-update.xlsm','Template','en_AE',en)
build(SRC+'CONSUMER_ELECTRONICS (1).xlsm',OUT+'NasCart-Amazon-AR-update.xlsm','النموذج','ar_AE',ar)
